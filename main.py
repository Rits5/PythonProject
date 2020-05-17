import openpyxl
import numpy as np
from openpyxl.chart import (
    LineChart,
    Reference,
)
from openpyxl import Workbook
import pandas as pd

file = open("trial1csv.txt", "r")
#write_f = open("write_file.txt", "w")
arr = file.read().split()

for i in range(0, len(arr)):
    arr[i] = arr[i].split(',')

result = np.array(arr, dtype=float)

print(result)
print(arr)
print(result[0][1])
print(type(result))

excel_file = 'test_file.xlsx'
wb = openpyxl.Workbook()

ws = wb.active
ws.title = 'testing'

for row in range(1, len(result)+1):
    for col in range(1, 3):
        ws.cell(row=row, column=col).value = result[row-1][col-1]

ws.column_dimensions['A'].width = ws.column_dimensions['B'].width = 15

graph = LineChart()
graph.title = 'Values'
graph.style = 13
graph.y_axis.title = 'y-axis'
graph.x_axis.title = 'x-axis'

data = Reference(ws, min_col=1, min_row=1, max_col=2, max_row=122)
graph.add_data(data, titles_from_data=True)

line = graph.series[0]
line.smooth = True

ws.add_chart(graph, 'D2')

wb.save(filename=excel_file)
file.close()

class Graph(LineChart):
    def __init__(self, graph_type_, title_, style_, y_axis_, x_axis_):
        self.type_ = graph_type_
        self.title_ = title_
        self.style_ = style_
        self.y_axis_ = y_axis_
        self.x_axis_ = x_axis_

    def graph_data(self):
        self.type = graph_type
        self.title_ = title_
        self.style_ = style_
        self.y_axis_ = y_axis_
        self.x_axis_ = x_axis_
